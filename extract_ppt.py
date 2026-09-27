from pptx import Presentation
import sys

prs = Presentation(r'ppt\BhuSure_SIH26018_Idea_Presentation.pptx')
for i, slide in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} ===')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    # Handle unicode characters
                    safe_text = text.encode('ascii', 'replace').decode('ascii')
                    print(f'  {safe_text}')
    print()