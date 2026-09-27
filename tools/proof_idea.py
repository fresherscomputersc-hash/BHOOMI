import unicodedata
from pptx import Presentation
o = Presentation("ppt/BhuSure_SIH26018_Idea_Presentation.pptx")
out = []
for i in range(6):
    out.append("#" * 15 + f" SLIDE {i + 1}")
    for sh in o.slides[i].shapes:
        if sh.has_text_frame and sh.name in ("Title 1", "Title 7", "TextBox 8", "TextBox 9") and sh.text_frame.text.strip():
            t = sh.text_frame.text.encode("ascii", "replace").decode()
            out.append(t[:900])
            out.append("")
open("tools/idea_proof.txt", "w", encoding="utf-8").write("\n".join(out))
print("written", len("\n".join(out)), "chars")
