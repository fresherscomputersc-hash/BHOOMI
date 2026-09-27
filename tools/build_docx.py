#!/usr/bin/env python3
"""Generate BhuSure full project documentation as a Word file.

Facts sourced from the repo (README, code, measured runs). No invented data.
Output: BhuSure_Full_Documentation.docx (repo root).
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor

NAVY = RGBColor(0x0A, 0x3D, 0x91)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def h1(t): doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)
def h3(t): doc.add_heading(t, level=3)
def p(t, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(t)
    run.bold = bold
    return para
def bullets(items):
    for b in items:
        doc.add_paragraph(b, style="List Bullet")
def table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    for j, htxt in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = ""
        r = c.paragraphs[0].add_run(htxt)
        r.bold = True
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            t.rows[i].cells[j].text = str(val)
    doc.add_paragraph()

# ---------------- Cover ----------------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("BhuSure")
r.font.size = Pt(36)
r.bold = True
r.font.color.rgb = NAVY
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Intelligent Land Record Digitization and Validation System")
r.font.size = Pt(16)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Smart India Hackathon 2026 \u2022 Problem Statement SIH26018\n"
             "Ministry of Rural Development \u2022 Theme: Smart Automation\n"
             "Team: Shadow Slayers \u2022 Full Project Documentation").font.size = Pt(12)
doc.add_page_break()

# ---------------- 1 Overview ----------------
h1("1. Overview")
p("BhuSure converts legacy land records \u2014 scanned registers, handwritten pages, "
  "PDFs and cadastral maps \u2014 into structured, validated, GIS-linked digital records. "
  "It is not an OCR wrapper: computer-vision preprocessing, printed OCR and handwritten "
  "recognition, structured field extraction with field-level confidence, ten named business "
  "rules, duplicate detection, cross-database checks and cadastral linking run first; then "
  "the result is handed to a government reviewer who holds final authority. Every AI output "
  "and every human correction is written to an immutable, SHA-256 hash-chained audit trail.")
p("Legal position: BhuSure is decision support. It does not adjudicate disputes, does not "
  "change ownership, and does not write back to Bhulekh, BhuNaksha, IGR or LRMS. "
  "Discrepancies are flagged, not resolved.", bold=True)

h2("1.1 Problem statement (SIH26018) requirements coverage")
table(["#", "Requirement", "Status"],
      [["7", "Multilingual recognition (major Indian languages)", "English/Hindi/Odia + script routing (~60%)"],
       ["8", "Extraction from PDFs, images, historical documents", "PDFs + images + degraded handling (~75%)"],
       ["9", "Classification into predefined fields", "17 fields, 6 doc profiles, anti-hallucination gates (~75%)"],
       ["10", "Business rules + cross-db + duplicate detection", "BR-1\u2026BR-10 tested; govt links mocked (~70%)"],
       ["11", "Confidence scoring + uncertain-field flagging", "Fusion scoring, bands, auto-review routing (~90%)"],
       ["12", "Human-assisted verification workflow", "Queue \u2192 correct \u2192 approve/block \u2192 audit (~90%)"],
       ["13", "AI learning over time", "Corrections captured, not yet trained on (~30%)"],
       ["14", "LRMS/DILRMP/GIS integration", "Live-ready contracts, mock data (~35%)"],
       ["15", "Secure repository + metadata + audit", "Hash chain, RBAC, argon2 (~70%)"],
       ["\u2014", "Interactive dashboards (7 asked panels)", "Live with real numbers (~90%)"],
       ["\u2014", "Government APIs", "FastAPI + interactive docs (~85%)"]])

# ---------------- 2 Architecture ----------------
h1("2. System architecture")
bullets([
    "Document upload (FR-1): PDF/JPG/PNG/TIFF, 25 MB cap, SHA-256 duplicate skip, batch tracking.",
    "CV preprocessing + layout (FR-2): deskew, denoise, CLAHE, adaptive binarisation, table/handwriting/stamp/map regions (OpenCV).",
    "OCR/HTR + word confidence (FR-3): Tesseract 5.5 (eng+hin+ori); script-routed second pass (39-A\u2192Oriya, Hindi belt\u2192Hindi, weak pages by decisive vote).",
    "Structured extraction (FR-4): 17 canonical fields; label aliases in Latin, Devanagari and Odia; regex validators; NER-style scoring; table-row/cell and pattern-sweep fallbacks.",
    "Confidence fusion (FR-5): 0.45\u00b7OCR + 0.30\u00b7pattern + 0.15\u00b7proximity + 0.10\u00b7consistency; bands \u226590 high, 70\u201389 medium, <70 low \u2192 mandatory review.",
    "Optional verifiers (pilot, flag-gated): Groq LLM (gpt-oss-120b), spaCy NER, HF transformers (XLM-R NER, TrOCR rescan). Never overwrite trusted values; never fill prohibited fields.",
    "GIS linking (FR-9): Shapely polygon areas on a sample cadastral layer.",
    "Validation BR-1\u2026BR-10 (FR-6/7) + cross-database checks (FR-8, mocked adapters with live-ready contracts).",
    "Reviewer verification (FR-10): approve/reject/escalate; approval server-refused (HTTP 409) with open critical findings.",
    "Audit trail (FR-12): SHA-256 hash chain with tamper detection.",
    "MIS dashboard, search, APIs (FR-11/13/17): measured accuracy, confidence bands, error fields, state-wise progress, CSV export.",
])

h2("2.1 Technology stack")
table(["Layer", "Prototype", "EC2 pilot", "Production path"],
      [["Web/API", "FastAPI, uvicorn", "same", "same + API keys/rate limits"],
       ["Persistence", "SQLite", "SQLite", "PostgreSQL + PostGIS"],
       ["Queue", "In-process thread worker", "same", "Redis/Celery"],
       ["OCR/HTR", "Tesseract 5.5 + OpenCV", "+ script routing", "TrOCR/IndicHTR (WordToken contract kept)"],
       ["NLP", " rapidfuzz + rules", "+ Groq/spaCy/XLM-R verifiers", "Fine-tuned per-state models"],
       ["Layout", "OpenCV heuristics", "same (+table-handwriting fix)", "YOLO/Detectron2 (same labelled boxes)"],
       ["GIS", "Shapely + sample GeoJSON", "same", "GeoServer + PostGIS"],
       ["External systems", "Deterministic mocks", "same", "Read-only govt APIs (ExternalCheck contract kept)"],
       ["Auth", "Seeded accounts + argon2", "same", "Department SSO + MFA"],
       ["Frontend", "Vanilla JS, no build, inline SVG maps", "same", "same"]])

h2("2.2 Document profiles")
table(["Profile", "Covers", "Key rule"],
      [["generic", "Any RoR/register page", "Standard required set"],
       ["odisha_khatiyan_39a", "Odisha Schedule I Form 39-A", "Khasra/khata/survey prohibited; parcel table is the area source; first khatiyan section only"],
       ["up_khatauni", "UP Khatauni (Gata)", "Gata\u2192plot_no; khasra/survey prohibited"],
       ["mp_khasra", "MP Khasra Panchshala", "Survey/plot prohibited"],
       ["bihar_khatiyan", "Bihar Jamabandi (Anchal)", "Survey/plot prohibited; \u0905\u0902\u091a\u0932 alias"],
       ["rajasthan_jamabandi", "Rajasthan Jamabandi", "Survey/plot prohibited"]])
p("Detection is heading-based with state gating plus rapidfuzz fallback for OCR-mangled headings.")

h2("2.3 Business rules BR-1\u2026BR-10")
bullets([
    "BR-1 mandatory fields (per document type) \u2014 critical/high.",
    "BR-2 sub-plot area reconciliation (skipped for identifier-free profiles).",
    "BR-3 duplicate khasra/owner conflict \u2014 critical, blocks approval.",
    "BR-4 chronological validity; BR-5 mutation-chain continuity.",
    "BR-6 unit normalisation (acre/hectare/decimal/bigha/katha/guntha/kanal/marla + Indic units).",
    "BR-7 village\u2013tehsil master check (scope-aware: out-of-scope districts route to review, low).",
    "BR-8 owner fuzzy match (rapidfuzz, cross-script aware roadmap).",
    "BR-9 GIS area reconciliation (±15% tolerance). BR-10 registration reference format.",
])

# ---------------- 3 Measured results ----------------
h1("3. Measured results (evidence, not claims)")
table(["Metric", "Figure", "Evidence"],
      [["Test suite", "165 local / 168 EC2, green", "pytest tests/"],
       ["Documents exercised", "6 demo + 5 stress + 5 Hindi + 5 English + 5 real Bhulekh RoRs", "report_full.json sets"],
       ["Determinism", "15/15 identical fields across reruns", "stability harness"],
       ["Adversarial acceptance", "36.8% field accuracy; 100% review-routed, 0 crashes, 0 hallucinated identifiers", "tools/acceptance.py"],
       ["Real Khatiyans", "360\u2013510 words, 79\u201384% OCR; owners/relations/parcel areas correct", "EC2 pilot records"],
       ["Hindi profiles", "up_khatauni/mp_khasra/bihar_khatiyan fire incl. fuzzy headings", "EC2 pilot records"],
       ["OCR stack", "Tesseract 5.5.0, eng+hin+ori+osd", "/api/v1/system/status"],
       ["Pipeline latency", "15\u2013127 s/doc on 2 vCPU/1 GB; validation < 1 s/record", "MIS performance panel"]])

# ---------------- 4 Deployment ----------------
h1("4. Deployment")
h2("4.1 Tiers")
table(["Tier", "Where", "Runs", "Notes"],
      [["Prototype/demo", "Render (Docker, free)", "main branch", "Sleeps when idle; SQLite resets on restart \u2014 demo only"],
       ["Live", "EC2 :8000 (bhusure.service)", "main branch via auto-deploy", "Real users; Groq-free path"],
       ["Pilot (full AI)", "EC2 :8001 (bhusure-pilot.service)", "pilot tree", "Groq + spaCy + XLM-R; TrOCR parked (Latin-only model)"]])
h2("4.2 Pilot enablement")
bullets([
    "SPACY_ENABLED=1 (en_core_web_sm), HF_ENABLED=1 (XLM-R NER), HF_HTR_ENABLED=0 (parked), GROQ_ENABLED=1 + GROQ_API_KEY, GROQ_MODEL=openai/gpt-oss-120b.",
    "OMP_THREAD_LIMIT=1 for reproducible OCR; 2 GB swap on small boxes (persisted in fstab).",
    "systemd units restart always; health endpoint /api/v1/health.",
])
h2("4.3 Seeded accounts")
table(["Username", "Password", "Role"],
      [["operator", "operator123", "Digitization Operator"],
       ["reviewer", "reviewer123", "Revenue Inspector"],
       ["supervisor", "supervisor123", "Supervisor (dashboards, export)"],
       ["survey", "survey123", "Survey Official"],
       ["auditor", "auditor123", "Auditor"],
       ["admin", "admin123", "Administrator"]])

# ---------------- 5 API ----------------
h1("5. API surface (FR-17)")
bullets([
    "POST /api/v1/auth/login \u2014 seeded SSO stand-in (token).",
    "POST /api/v1/documents (multipart, process_sync flag) \u2014 bulk upload, dedup by hash.",
    "POST /api/v1/documents/{id}/process \u2014 reprocess; GET detail incl. OCR, record, audit.",
    "GET /api/v1/records/{id} (+/discrepancies) \u2014 record, field confidences/sources, findings.",
    "POST /api/v1/records/{id}/verify (approve/reject/escalate) \u2014 409 on open criticals.",
    "GET /api/v1/dashboard (+/alerts, /export.csv) \u2014 MIS numbers; GET /api/v1/system/status.",
    "Interactive docs at /docs. RBAC enforced per route (operator cannot read audit).",
])

# ---------------- 6 Testing ----------------
h1("6. Testing and quality gates")
bullets([
    "165 tests locally (168 on EC2 where Tesseract executes): units, table/formats, all BR-1\u2026BR-10 incl. must-not-fire cases, deskew proof, GIS reconciliation, audit tamper detection, API acceptance, RBAC.",
    "Regression tests for every production incident (table date-tuple crash, header leaks, BR-2 empty guard, HTR clamp, routing gates).",
    "Acceptance harness (tools/acceptance.py) scores records against per-set ground truth; stability harness proves rerun determinism.",
    "Every fix ships with a test; suites run on localhost and EC2 before deploy.",
])

# ---------------- 7 Security & legal ----------------
h1("7. Security and legal position")
bullets([
    "argon2 password hashing; opaque bearer tokens; RBAC matrix on every route.",
    "Append-only SHA-256 audit chain; tamper re-verification endpoint.",
    "PII discipline: real owner data never leaves the box for external APIs (Groq disabled on real-data runs); secrets in 600-perm env files, never in git.",
    "Legal: decision support only \u2014 no adjudication, no ownership changes, no write-back without government approval.",
])

# ---------------- 8 Roadmap ----------------
h1("8. Roadmap to production")
table(["Gap", "Achieve by", "Needs"],
      [["Live govt integrations", "Swap mock adapters (contracts match)", "Department API access"],
       ["Real learning loop", "Fine-tune NER/HTR on correction dataset", "GPU node + correction volume"],
       ["More languages", "Profile packs (Bengali/Tamil/Telugu/Marathi)", "Packs + samples"],
       ["Indic handwriting", "IndicHTR fine-tune (TrOCR-base is Latin-only)", "GPU + labelled pages"],
       ["Hardening", "Postgres HA, Redis/Celery, K8s, WORM audit, SSO/MFA, backups, DR", "Production funding"]])

doc.add_page_break()
p("Generated from the BhuSure repository state. Figures above are measured on the EC2 pilot unless noted.",
  bold=True)

doc.save("BhuSure_Full_Documentation.docx")
print("wrote BhuSure_Full_Documentation.docx")
